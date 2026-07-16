from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from hairgpu_final_db import FinalStageDatabase

HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
PASS_FILL = PatternFill(fill_type="solid", fgColor="C6EFCE")
WARN_FILL = PatternFill(fill_type="solid", fgColor="FFEB9C")
FAIL_FILL = PatternFill(fill_type="solid", fgColor="FFC7CE")


def _row_dicts(rows: Iterable[Any]) -> list[dict[str, Any]]:
    return [dict(row) for row in rows]


def _append_table(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["status", "no data"])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        ws.append([row.get(header) for header in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _autofit(ws) -> None:
    for column_index, column in enumerate(ws.columns, 1):
        width = max(
            (len(str(cell.value)) if cell.value is not None else 0)
            for cell in column
        )
        ws.column_dimensions[get_column_letter(column_index)].width = min(
            max(width + 2, 12), 52
        )
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)


def _comparison_rows(db: FinalStageDatabase, method: str, baseline: str) -> list[dict[str, Any]]:
    rows = _row_dicts(db.fetch_comparison_rows())
    lookup = {(row["scenario_code"], row["method_code"]): row for row in rows}
    scenarios = sorted({row["scenario_code"] for row in rows})
    metrics = (
        "avg_frame_ms",
        "avg_p95_frame_ms",
        "avg_p99_frame_ms",
        "avg_hair_gpu_ms",
        "avg_sixty_fps_ratio",
        "avg_false_static_rate",
        "avg_predictive_wake_recall",
        "avg_repair_verified_success_rate",
        "avg_collision_miss_rate",
        "avg_state_oscillation_rate",
        "avg_max_boundary_error_mm",
    )
    output: list[dict[str, Any]] = []
    for scenario in scenarios:
        target = lookup.get((scenario, method))
        reference = lookup.get((scenario, baseline))
        if not target or not reference:
            continue
        row: dict[str, Any] = {
            "scenario_code": scenario,
            "method_code": method,
            "baseline_method_code": baseline,
        }
        for metric in metrics:
            target_value = target.get(metric)
            baseline_value = reference.get(metric)
            row[f"method_{metric}"] = target_value
            row[f"baseline_{metric}"] = baseline_value
            row[f"delta_{metric}"] = (
                float(target_value) - float(baseline_value)
                if target_value is not None and baseline_value is not None
                else None
            )
        output.append(row)
    return output


def _formal_checks(run_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    checks: list[dict[str, Any]] = []
    for row in run_rows:
        identity_match = len(
            {
                str(row.get("method_code") or ""),
                str(row.get("requested_method_code") or ""),
                str(row.get("applied_method_code") or ""),
                str(row.get("observed_method_code") or ""),
            }
        ) == 1
        checks.append(
            {
                "run_name": row.get("run_name"),
                "method_identity": "PASS" if identity_match else "FAIL",
                "validation_status": row.get("validation_status"),
                "real_measurement": "PASS" if row.get("real_measurement") == 1 else "FAIL",
                "thermal_valid": "PASS" if row.get("thermal_valid") in (None, 1) else "FAIL",
                "environment_valid": "PASS" if row.get("environment_valid") in (None, 1) else "FAIL",
                "metric_coverage_valid": "PASS" if row.get("required_metric_coverage_valid") in (None, 1) else "FAIL",
                "manifest_present": "PASS" if row.get("manifest_json_path") else "FAIL",
                "canonical_hash_present": "PASS" if row.get("canonical_result_sha256") else "FAIL",
                "result_manifest_hash_present": "PASS" if row.get("result_manifest_sha256") else "FAIL",
                "gpu_timing_coverage": row.get("gpu_timing_coverage"),
                "same_frame_readback_count": row.get("same_frame_readback_count"),
                "synthetic_metric_count": row.get("synthetic_metric_count"),
                "fatal_runtime_failure_count": row.get("fatal_runtime_failure_count"),
            }
        )
    return checks


def export_workbook(db_path: Path, output_xlsx: Path) -> Path:
    db = FinalStageDatabase(db_path)
    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "FormalRunSummary"
    summary = db.fetch_experiment_summary()
    _append_table(
        summary_sheet,
        [{"metric": key, "value": value} for key, value in summary.items()],
    )

    run_rows = _row_dicts(db.fetch_runs())
    frame_rows = _row_dicts(db.fetch_frame_metrics())
    method_rows = db.fetch_method_summary_rows()
    scenario_rows = db.fetch_scenario_summary_rows()

    sheets: list[tuple[str, list[dict[str, Any]]]] = [
        ("MethodComparison", method_rows),
        ("ScenarioSummary", scenario_rows),
        ("B8_Representation", _comparison_rows(db, "P", "B8")),
        ("B6_Performance", _comparison_rows(db, "P", "B6")),
        ("B0_QualityReference", _comparison_rows(db, "P", "B0")),
        ("FormalValidity", _formal_checks(run_rows)),
        ("RawRuns", run_rows),
        ("RawFrames", frame_rows),
        ("Artifacts", db.fetch_artifact_rows()),
        ("DerivedMetrics", db.fetch_derived_metric_rows()),
        ("ScenarioCoverage", db.fetch_scenario_coverage_rows()),
        ("HypothesisChecks", db.fetch_hypothesis_rows()),
        ("UserTasks", db.fetch_user_task_rows()),
        ("UserTaskSummary", db.fetch_user_task_summary_rows()),
    ]

    quality_columns = {
        "FalseStaticAnalysis": (
            "run_name", "method_code", "scenario_code", "false_static_rate"
        ),
        "WakeAnalysis": (
            "run_name", "method_code", "scenario_code",
            "predictive_wake_precision", "predictive_wake_recall"
        ),
        "RepairAnalysis": (
            "run_name", "method_code", "scenario_code",
            "repair_execution_success_rate", "repair_verified_success_rate",
            "repair_recurrence_rate"
        ),
        "BoundaryAnalysis": (
            "run_name", "method_code", "scenario_code", "max_boundary_error_mm"
        ),
        "GpuPassTiming": (
            "run_name", "method_code", "scenario_code", "avg_hair_gpu_ms",
            "p95_hair_gpu_ms", "p99_hair_gpu_ms", "gpu_timing_coverage"
        ),
        "HardwareEnvironment": (
            "run_name", "hardware_profile", "gpu_name", "gpu_driver",
            "dedicated_vram_mb", "cpu_name", "power_scheme",
            "thermal_valid", "environment_valid"
        ),
        "Determinism": (
            "run_name", "method_code", "scenario_code", "seed",
            "decision_sequence_sha256", "state_sequence_sha256",
            "compaction_sequence_sha256", "final_position_sha256"
        ),
        "ArtifactIntegrity": (
            "run_name", "manifest_json_path", "result_manifest_sha256",
            "canonical_result_sha256", "summary_sha256", "frame_csv_sha256",
            "frame_jsonl_sha256", "settings_json_sha256",
            "environment_json_sha256", "validation_json_sha256"
        ),
    }
    for name, columns in quality_columns.items():
        rows = [{column: row.get(column) for column in columns} for row in run_rows]
        sheets.append((name, rows))

    invalid_rows = [
        row for row in run_rows
        if row.get("validation_status") != "complete_valid" or row.get("is_valid") != 1
    ]
    sheets.append(("InvalidRuns", invalid_rows))

    for name, rows in sheets:
        sheet = workbook.create_sheet(name[:31])
        _append_table(sheet, rows)
        if sheet.max_row >= 2 and sheet.max_column >= 1:
            for column in range(1, sheet.max_column + 1):
                header = sheet.cell(1, column).value
                if header and (
                    str(header).endswith("_valid")
                    or str(header).endswith("_present")
                    or str(header) in {"method_identity", "real_measurement"}
                ):
                    column_letter = get_column_letter(column)
                    sheet.conditional_formatting.add(
                        f"{column_letter}2:{column_letter}{sheet.max_row}",
                        CellIsRule(operator="equal", formula=['"PASS"'], fill=PASS_FILL),
                    )
                    sheet.conditional_formatting.add(
                        f"{column_letter}2:{column_letter}{sheet.max_row}",
                        CellIsRule(operator="equal", formula=['"FAIL"'], fill=FAIL_FILL),
                    )

    for sheet in workbook.worksheets:
        _autofit(sheet)

    output_xlsx = Path(output_xlsx)
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_xlsx)
    db.close()
    return output_xlsx
