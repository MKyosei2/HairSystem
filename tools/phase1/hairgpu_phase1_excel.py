from __future__ import annotations

import sqlite3
from pathlib import Path
from typing import Iterable, Sequence

import openpyxl
from openpyxl.styles import Font

from hairgpu_phase1_db import fetch_ablation_rows, fetch_comparison_rows


def _write_sheet(workbook: openpyxl.Workbook, name: str, headers: Sequence[str], rows: Iterable) -> None:
    ws = workbook.create_sheet(title=name[:31])
    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        if isinstance(row, sqlite3.Row):
            ws.append([row[h] for h in headers])
        elif isinstance(row, dict):
            ws.append([row.get(h) for h in headers])
        else:
            ws.append(list(row))
    ws.freeze_panes = "A2"
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 12), 42)


def export_workbook(conn: sqlite3.Connection, output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    default = wb.active
    wb.remove(default)

    for table_name in [
        "experiments",
        "scenarios",
        "methods",
        "runs",
        "metrics_summary",
        "metrics_frame",
        "teacher_datasets",
        "teacher_frame_summary",
        "learned_models",
        "active_set_evaluations",
        "corrective_physics_evaluations",
        "external_contact_evaluations",
    ]:
        rows = conn.execute(f"SELECT * FROM {table_name}").fetchall()
        headers = list(rows[0].keys()) if rows else []
        _write_sheet(wb, table_name, headers, rows)

    comparison_rows = fetch_comparison_rows(conn)
    comparison_headers = list(comparison_rows[0].keys()) if comparison_rows else [
        "experiment_name", "method_code", "method_name", "scenario_code", "scenario_name", "asset_name",
        "runtime_mode", "frame_count", "target_fps", "frame_time_p95_ms", "worst_frame_time_ms", "peak_ram_mb",
        "penetration_count_per_sec", "max_penetration_depth_mm", "silhouette_error", "clump_preservation_score",
        "edit_to_run_latency_ms", "prior_rmse", "residual_rmse", "full_model_rmse", "uncertainty_mae",
        "average_selected_cluster_count", "average_selected_guide_count", "average_budget_utilization",
        "penetration_reduction_ratio", "silhouette_recovery_ratio", "correction_cost_ms",
        "total_contact_event_count", "contact_penetration_reduction_ratio", "valid",
        "validation_warning_count", "validation_error_count",
    ]
    _write_sheet(wb, "phase7_comparison", comparison_headers, comparison_rows)

    ablation_rows = fetch_ablation_rows(conn)
    ablation_headers = list(ablation_rows[0].keys()) if ablation_rows else [
        "experiment_name", "scenario_code", "asset_name", "method_code", "delta_p95_ms_vs_A",
        "delta_penetration_mm_vs_A", "delta_silhouette_error_vs_A", "delta_clump_score_vs_A",
        "prior_rmse", "residual_rmse", "penetration_reduction_ratio", "average_selected_cluster_count",
    ]
    _write_sheet(wb, "phase7_ablation", ablation_headers, ablation_rows)

    wb.save(output_path)
    return output_path
