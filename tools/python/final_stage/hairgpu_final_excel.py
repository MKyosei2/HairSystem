from __future__ import annotations

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from hairgpu_final_db import FinalStageDatabase

HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _style_header(row):
    for cell in row:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT


def export_workbook(db_path: Path, output_xlsx: Path) -> Path:
    db = FinalStageDatabase(db_path)
    wb = Workbook()
    ws_runs = wb.active
    ws_runs.title = "Runs"
    run_rows = db.fetch_runs()
    if run_rows:
        headers = list(run_rows[0].keys())
        ws_runs.append(headers)
        _style_header(ws_runs[1])
        for row in run_rows:
            ws_runs.append([row[h] for h in headers])

    ws_frames = wb.create_sheet("FrameMetrics")
    frame_rows = db.fetch_frame_metrics()
    if frame_rows:
        headers = list(frame_rows[0].keys())
        ws_frames.append(headers)
        _style_header(ws_frames[1])
        for row in frame_rows:
            ws_frames.append([row[h] for h in headers])

    ws_comp = wb.create_sheet("Comparison")
    comp_rows = db.fetch_comparison_rows()
    if comp_rows:
        headers = list(comp_rows[0].keys())
        ws_comp.append(headers)
        _style_header(ws_comp[1])
        for row in comp_rows:
            ws_comp.append([row[h] for h in headers])

    ws_ab = wb.create_sheet("Ablation")
    ab_rows = db.fetch_ablation_rows()
    if ab_rows:
        headers = list(ab_rows[0].keys())
        ws_ab.append(headers)
        _style_header(ws_ab[1])
        for row in ab_rows:
            ws_ab.append([row[h] for h in headers])

    for ws in wb.worksheets:
        for column in ws.columns:
            width = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column) + 2
            ws.column_dimensions[column[0].column_letter].width = min(max(width, 12), 36)

    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_xlsx)
    db.close()
    return output_xlsx
